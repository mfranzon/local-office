"""Turns uploaded files into searchable pages: render -> (text layer | OCR) -> chunk -> embed -> index.

One background worker processes documents one page at a time, so progress is visible and each
finished page is searchable immediately.
"""
import asyncio
import io
import logging
import re
import shutil
import uuid
from collections.abc import Iterator
from html.parser import HTMLParser
from pathlib import Path
from typing import BinaryIO

import pypdfium2 as pdfium
from docx import Document as Docx
from openpyxl import load_workbook
from PIL import Image, ImageOps
from pillow_heif import register_heif_opener

from . import llm, models, wiki
from .config import DATA_DIR
from .db import db, vec

register_heif_opener()
log = logging.getLogger(__name__)

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".heic", ".heif", ".tif", ".tiff", ".webp", ".bmp"}
TEXT_EXTS = {".txt", ".md", ".csv"}
SUPPORTED = IMAGE_EXTS | TEXT_EXTS | {".pdf", ".docx", ".xlsx"}
OCR_SIDE = 1540            # LightOnOCR's recommended longest image side
MIN_TEXT_LAYER = 80        # fewer chars than this on a PDF page means it's a scan -> OCR
CHUNK_CHARS = 1200

FILES = DATA_DIR / "files"
PAGES = DATA_DIR / "pages"
_queue: asyncio.Queue[str] = asyncio.Queue()


def file_path(doc_id: str, ext: str) -> Path:
    return FILES / f"{doc_id}{ext}"


def page_image_path(doc_id: str, n: int) -> Path:
    return PAGES / doc_id / f"{n}.jpg"


# --- page sources: each yields (image or None, text or None) per page ---

def _fit(img: Image.Image) -> Image.Image:
    img = ImageOps.exif_transpose(img).convert("RGB")
    img.thumbnail((OCR_SIDE, OCR_SIDE), Image.LANCZOS)
    return img


def _pdf_pages(path: Path) -> Iterator[tuple[Image.Image, str | None]]:
    pdf = pdfium.PdfDocument(path)
    try:
        for page in pdf:
            text = page.get_textpage().get_text_range()
            scale = min(3.0, OCR_SIDE / max(page.get_size()))
            img = page.render(scale=scale).to_pil().convert("RGB")
            yield img, text if len(text.strip()) >= MIN_TEXT_LAYER else None
    finally:
        pdf.close()


def _image_pages(path: Path) -> Iterator[tuple[Image.Image, None]]:
    with Image.open(path) as im:   # multi-page TIFFs from scanners have several frames
        for i in range(getattr(im, "n_frames", 1)):
            im.seek(i)
            yield _fit(im.copy()), None


def _docx_pages(path: Path) -> Iterator[tuple[None, str]]:
    doc = Docx(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts += [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
    yield None, "\n".join(parts)


def _xlsx_pages(path: Path) -> Iterator[tuple[None, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ("\t".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True))
        yield None, f"# {ws.title}\n" + "\n".join(r for r in rows if r.strip())


def _text_pages(path: Path) -> Iterator[tuple[None, str]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    for i in range(0, max(len(text), 1), 4000):
        yield None, text[i:i + 4000]


def _pages(path: Path) -> Iterator[tuple[Image.Image | None, str | None]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf_pages(path)
    if ext in IMAGE_EXTS:
        return _image_pages(path)
    return {".docx": _docx_pages, ".xlsx": _xlsx_pages}.get(ext, _text_pages)(path)


def _count(path: Path) -> int:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return len(pdfium.PdfDocument(path))
    if ext in IMAGE_EXTS:
        with Image.open(path) as im:
            return getattr(im, "n_frames", 1)
    return sum(1 for _ in _pages(path))


def chunk(text: str) -> list[str]:
    """Paragraph-aware chunks of ~CHUNK_CHARS, carrying the last paragraph over for context."""
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks, cur = [], []
    for p in paras:
        if cur and sum(map(len, cur)) + len(p) > CHUNK_CHARS:
            chunks.append("\n\n".join(cur))
            cur = cur[-1:] if len(cur[-1]) < CHUNK_CHARS // 3 else []
        cur.append(p)
    if cur:
        chunks.append("\n\n".join(cur))
    return [c[i:i + CHUNK_CHARS * 2] for c in chunks for i in range(0, len(c), CHUNK_CHARS * 2)]


# --- worker ---

class _TableToMarkdown(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows: list[list[str]] = []
        self.cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.rows.append([])
        elif tag in ("td", "th"):
            self.cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.cell is not None and self.rows:
            self.rows[-1].append(" ".join("".join(self.cell).split()).replace("|", "\\|"))
            self.cell = None

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)


def html_tables_to_markdown(text: str) -> str:
    """LightOnOCR emits tables as HTML; pipe tables read better and embed with less noise."""
    def convert(m: re.Match) -> str:
        p = _TableToMarkdown()
        p.feed(m.group(0))
        rows = [r for r in p.rows if r]
        if not rows:
            return ""
        width = max(map(len, rows))
        rows = [r + [""] * (width - len(r)) for r in rows]
        lines = ["| " + " | ".join(r) + " |" for r in rows]
        lines.insert(1, "|" + " --- |" * width)
        return "\n".join(lines)
    return re.sub(r"<table.*?</table>", convert, text, flags=re.S | re.I)


async def _ocr(img: Image.Image) -> str:
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=90)
    content = [llm.image_part(buf.getvalue()), {"type": "text", "text": models.active("ocr").prompt}]
    return html_tables_to_markdown(await llm.chat("ocr", [{"role": "user", "content": content}], max_tokens=4096))


async def _index_page(doc_id: str, name: str, n: int, text: str):
    pieces = chunk(text)
    async with models.index_lock:  # never write vectors while the index is being rebuilt
        vectors = await llm.embed([f"{name}\n{p}" for p in pieces]) if pieces else []
        _store_chunks(doc_id, n, pieces, vectors)


def _store_chunks(doc_id: str, n: int, pieces: list[str], vectors: list[list[float]]):
    with db() as conn:
        for piece, v in zip(pieces, vectors):
            cid = conn.execute("INSERT INTO chunks(doc_id, page, text) VALUES (?,?,?)", (doc_id, n, piece)).lastrowid
            conn.execute("INSERT INTO vec_chunks(rowid, embedding) VALUES (?,?)", (cid, vec(v)))


# --- adding and removing documents (uploads, synced folders) ---

def add(name: str, ext: str, *, data: bytes | None = None, src: Path | BinaryIO | None = None,
        source_id: int | None = None, source_key: str | None = None, mtime: float | None = None) -> str:
    """Stores a copy of the file and queues it for reading. Returns the new document id."""
    doc_id = uuid.uuid4().hex[:12]
    dest = file_path(doc_id, ext)
    if isinstance(src, Path):
        shutil.copyfile(src, dest)
    elif src is not None:
        with dest.open("wb") as out:
            shutil.copyfileobj(src, out)
    else:
        dest.write_bytes(data or b"")
    with db() as conn:
        conn.execute("""INSERT INTO documents(id, name, ext, size, source_id, source_key, source_mtime)
                        VALUES (?,?,?,?,?,?,?)""", (doc_id, name, ext, dest.stat().st_size, source_id, source_key, mtime))
    enqueue(doc_id)
    return doc_id


def replace(doc_id: str, ext: str, src: Path, mtime: float):
    """A synced file changed: read it again."""
    shutil.copyfile(src, file_path(doc_id, ext))
    with db() as conn:
        conn.execute("UPDATE documents SET size=?, source_mtime=?, status='queued' WHERE id=?",
                     (file_path(doc_id, ext).stat().st_size, mtime, doc_id))
    enqueue(doc_id)


def remove(conn, doc_ids: list[str]) -> list[int]:
    """Deletes documents with their files. Returns the wiki entities they cited."""
    cited = set()
    for doc_id in doc_ids:
        doc = conn.execute("SELECT name, ext FROM documents WHERE id=?", (doc_id,)).fetchone()
        if doc is None:
            continue
        cited |= {r[0] for r in conn.execute("SELECT DISTINCT entity_id FROM facts WHERE doc_id=? AND entity_id IS NOT NULL", (doc_id,))}
        conn.execute("DELETE FROM documents WHERE id=?", (doc_id,))   # its facts go with it (ON DELETE CASCADE)
        conn.execute("INSERT INTO wiki_log(action, detail) VALUES ('delete', ?)", (doc["name"],))
        file_path(doc_id, doc["ext"]).unlink(missing_ok=True)
        shutil.rmtree(PAGES / doc_id, ignore_errors=True)
    conn.executemany("UPDATE entities SET summary_stale=1 WHERE id=?", [(e,) for e in cited])
    wiki.cleanup(conn)
    return sorted(cited)


async def process(doc_id: str):
    with db() as conn:
        doc = conn.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
        if doc is None:
            return
        conn.execute("DELETE FROM chunks WHERE doc_id=?", (doc_id,))
        conn.execute("DELETE FROM pages WHERE doc_id=?", (doc_id,))
        conn.execute("UPDATE documents SET status='processing', pages_done=0, error=NULL WHERE id=?", (doc_id,))
    path = file_path(doc_id, doc["ext"])
    total = await asyncio.to_thread(_count, path)
    with db() as conn:
        conn.execute("UPDATE documents SET pages_total=? WHERE id=?", (total, doc_id))

    pages = _pages(path)
    (PAGES / doc_id).mkdir(parents=True, exist_ok=True)
    for n in range(1, total + 1):
        item = await asyncio.to_thread(next, pages, None)
        if item is None:
            break
        img, text = item
        if img is not None:
            await asyncio.to_thread(img.save, page_image_path(doc_id, n), "JPEG", quality=82)
        method = "text" if text is not None else "ocr"
        if text is None:
            text = await _ocr(img)
        with db() as conn:
            conn.execute("INSERT INTO pages(doc_id, n, text, method, has_image) VALUES (?,?,?,?,?)",
                         (doc_id, n, text, method, img is not None))
        await _index_page(doc_id, doc["name"], n, text)
        with db() as conn:
            conn.execute("UPDATE documents SET pages_done=? WHERE id=?", (n, doc_id))
    with db() as conn:
        conn.execute("UPDATE documents SET status='ready' WHERE id=?", (doc_id,))
    wiki.enqueue(doc_id)


async def worker():
    while True:
        doc_id = await _queue.get()
        try:
            await process(doc_id)
        except Exception as e:
            log.exception("ingest failed for %s", doc_id)
            with db() as conn:
                conn.execute("UPDATE documents SET status='error', error=? WHERE id=?", (str(e)[:500], doc_id))


def enqueue(doc_id: str):
    _queue.put_nowait(doc_id)


def resume_pending():
    with db() as conn:
        rows = conn.execute("SELECT id FROM documents WHERE status IN ('queued','processing') ORDER BY created_at")
        for r in rows:
            enqueue(r["id"])
