"""What a skill can do. Every tool only reads the office's documents and wiki, except create_file, which
writes a new file for the user to download. Nothing here sends, moves or deletes anything, or runs code.

Tool results number their sources ([1], [2], ...) across the whole run, so the final answer can cite them.
"""
import asyncio
import re
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path

from docx import Document as Docx
from openpyxl import Workbook

from . import rag, wiki_views
from .config import DATA_DIR
from .db import db
from .wiki import fold, name_norm

EXPORTS = DATA_DIR / "exports"
READ_CHARS = 8000

SPECS = [
    ("search_documents", "Search the office documents for passages about a topic. Returns numbered excerpts.",
     {"query": {"type": "string"}, "limit": {"type": "integer", "description": "default 6"}}, ["query"]),
    ("list_documents", "List documents, newest first, with their ids. Use it to find recent documents or a document by name.",
     {"days": {"type": "integer", "description": "only the last N days"},
      "text": {"type": "string", "description": "only names containing this"},
      "limit": {"type": "integer", "description": "default 20"}}, []),
    ("read_document", "Read the full text of a document by its id.", {"doc_id": {"type": "string"}}, ["doc_id"]),
    ("wiki_page", "The wiki page about an organization, person or project: details, links, deadlines and history.",
     {"name": {"type": "string"}}, ["name"]),
    ("list_deadlines", "Deadlines found in the documents (payments, deliveries, expiries, renewals).",
     {"days": {"type": "integer", "description": "how many days ahead, default 60"}}, []),
    ("create_file", "Create a file for the user to download: a Word document or Markdown from Markdown text, or an Excel sheet from rows.",
     {"filename": {"type": "string", "description": "without extension"},
      "format": {"type": "string", "enum": ["docx", "xlsx", "md"]},
      "content": {"type": "string", "description": "Markdown text, for docx and md"},
      "rows": {"type": "array", "items": {"type": "array", "items": {"type": ["string", "number", "null"]}},
               "description": "for xlsx: the first row is the header"}}, ["filename", "format"]),
]


def schemas(names: set[str] | None = None) -> list[dict]:
    return [{"type": "function", "function": {"name": n, "description": d, "parameters": {
        "type": "object", "properties": props, "required": req}}} for n, d, props, req in SPECS if names is None or n in names]


class Run:
    """State of one skill run: numbered sources and created files."""

    def __init__(self):
        self.sources: list[dict] = []
        self.files: list[dict] = []
        self.id = uuid.uuid4().hex[:10]

    def cite(self, **source) -> int:
        key = (source.get("kind", "doc"), source.get("doc_id"), source.get("entity_id"), source.get("page"))
        for s in self.sources:
            if (s["kind"], s["doc_id"], s["entity_id"], s["page"]) == key:
                return s["n"]
        n = len(self.sources) + 1
        self.sources.append({"n": n, "kind": "doc", "doc_id": None, "entity_id": None, "page": None, "text": "", **source})
        return n


async def search_documents(run: Run, query: str, limit: int = 6) -> str:
    hits = await rag.search(query, None, k=max(1, min(int(limit or 6), 12)))
    if not hits:
        return "No matching passages."
    return "\n\n".join(f"[{run.cite(doc_id=h['doc_id'], name=h['name'], page=h['page'], text=h['text'][:400])}] "
                       f"{h['name']}, p. {h['page']} (doc_id {h['doc_id']}):\n{h['text']}" for h in hits)


def list_documents(run: Run, days: int | None = None, text: str | None = None, limit: int = 20) -> str:
    with db() as conn:
        rows = [dict(r) for r in conn.execute("SELECT * FROM documents WHERE status='ready' ORDER BY created_at DESC")]
    if days:
        since = datetime.now() - timedelta(days=int(days))
        rows = [r for r in rows if datetime.fromisoformat(r["created_at"]) >= since]
    if text:
        t = fold(text)
        rows = [r for r in rows if t in fold(r["name"])]
    rows = rows[:max(1, min(int(limit or 20), 50))]
    return "\n".join(f"- doc_id {r['id']}: {r['name']}, added {r['created_at'][:10]}, {r['pages_total']} page(s)"
                     for r in rows) if rows else "No documents match."


def read_document(run: Run, doc_id: str) -> str:
    with db() as conn:
        doc = conn.execute("SELECT name FROM documents WHERE id=?", (doc_id,)).fetchone()
        pages = conn.execute("SELECT n, text FROM pages WHERE doc_id=? ORDER BY n", (doc_id,)).fetchall()
    if doc is None:
        return f"No document with id {doc_id}. Use list_documents or search_documents to find ids."
    out, used = [], 0
    for p in pages:
        n = run.cite(doc_id=doc_id, name=doc["name"], page=p["n"], text=p["text"][:400])
        text = p["text"][:READ_CHARS - used]
        out.append(f"[{n}] {doc['name']}, p. {p['n']}:\n{text}")
        used += len(text)
        if used >= READ_CHARS:
            out.append("(truncated)")
            break
    return "\n\n".join(out) or "This document has no text yet."


def wiki_page(run: Run, name: str) -> str:
    n = name_norm(name)
    with db() as conn:
        row = conn.execute("SELECT entity_id FROM aliases WHERE norm=?", (n,)).fetchone()
        if row is None:
            row = conn.execute("SELECT a.entity_id FROM aliases a WHERE a.norm LIKE ? LIMIT 1", (f"%{n}%",)).fetchone()
        if row is None:
            return f"No wiki page for {name}."
        eid = row[0]
        title = conn.execute("SELECT name FROM entities WHERE id=?", (eid,)).fetchone()[0]
        text = wiki_views.entity_markdown(conn, eid, for_llm=True)
    return f"[{run.cite(kind='wiki', entity_id=eid, name=title, text=text[:400])}] Wiki: {title}\n{text}"


def list_deadlines(run: Run, days: int = 60) -> str:
    until = (date.today() + timedelta(days=int(days or 60))).isoformat()
    with db() as conn:
        items = [d for d in wiki_views.deadlines(conn, past_days=14) if d["date"] <= until]
    if not items:
        return "No deadlines found."
    return "\n".join(f"[{run.cite(doc_id=d['doc_id'], name=d['doc_name'], page=d['page'], text=d['value'])}] {d['date']}"
                     f"{' (overdue)' if d['overdue'] else ''}: {d['value']} - {', '.join(e['name'] for e in d['entities']) or 'n/a'}"
                     for d in items)


def _safe_name(filename: str, ext: str) -> str:
    stem = re.sub(r"\.(docx|xlsx|md)$", "", filename.strip(), flags=re.I)
    stem = re.sub(r'[\\/:*?"<>|]+', "-", stem).strip(" .-") or "export"
    return f"{stem[:80]}.{ext}"


def markdown_to_docx(md: str, path: Path):
    doc = Docx()
    table: list[list[str]] = []

    def flush():
        if table:
            rows = [r for r in table if not all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in r)]
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    t.cell(i, j).text = c.strip()
            table.clear()

    for line in md.splitlines():
        s = line.strip()
        if s.startswith("|") and s.endswith("|"):
            table.append(s.strip("|").split("|"))
            continue
        flush()
        if m := re.match(r"(#{1,4})\s+(.*)", s):
            doc.add_heading(m[2], level=len(m[1]))
        elif m := re.match(r"[-*]\s+(.*)", s):
            _runs(doc.add_paragraph(style="List Bullet"), m[1])
        elif m := re.match(r"\d+[.)]\s+(.*)", s):
            _runs(doc.add_paragraph(style="List Number"), m[1])
        elif s:
            _runs(doc.add_paragraph(), s)
    flush()
    doc.save(path)


def _runs(paragraph, text: str):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):   # **bold**
        paragraph.add_run(part).bold = i % 2 == 1


def create_file(run: Run, filename: str, format: str, content: str | None = None, rows: list | None = None) -> str:
    if format not in ("docx", "xlsx", "md"):
        return "format must be docx, xlsx or md"
    folder = EXPORTS / run.id
    folder.mkdir(parents=True, exist_ok=True)
    name = _safe_name(filename, format)
    path = folder / name
    if format == "xlsx":
        if not rows:
            return "rows are required for xlsx"
        wb = Workbook()
        for r in rows:
            wb.active.append([c if isinstance(c, (int, float)) or c is None else str(c) for c in r])
        wb.save(path)
    elif not content:
        return "content is required for docx and md"
    elif format == "docx":
        markdown_to_docx(content, path)
    else:
        path.write_text(content, encoding="utf-8")
    run.files.append({"name": name, "url": f"/api/exports/{run.id}/{name}"})
    return (f"Created {name}. A download button appears under your answer automatically: do not write a link to it. "
            "Still give the full answer in the chat (e.g. the table), as the skill asks.")


FUNCTIONS = {"search_documents": search_documents, "list_documents": list_documents, "read_document": read_document,
             "wiki_page": wiki_page, "list_deadlines": list_deadlines, "create_file": create_file}


async def call(run: Run, name: str, args: dict) -> str:
    fn = FUNCTIONS.get(name)
    if fn is None:
        return f"Unknown tool {name}."
    allowed = {p for n, _, props, _ in SPECS if n == name for p in props}
    try:
        result = fn(run, **{k: v for k, v in args.items() if k in allowed})
        return await result if asyncio.iscoroutine(result) else result
    except Exception as e:   # the model sees the problem and can try something else
        return f"Error: {e}"
