import asyncio
import io
import json
import zipfile

import pytest
from docx import Document as Docx
from openpyxl import load_workbook

from app import agent, config, ingest, llm, skills, tools
from app import db as dbmod

SKILL = """---
name: weekly-brief
description: Summarize the week's documents.
---

1. Call list_documents with days 7.
"""


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(dbmod, "DATA_DIR", tmp_path)
    monkeypatch.setattr(ingest, "FILES", tmp_path / "files")
    monkeypatch.setattr(ingest, "enqueue", lambda doc_id: None)
    monkeypatch.setattr(skills, "USER_SKILLS", tmp_path / "skills")
    monkeypatch.setattr(tools, "EXPORTS", tmp_path / "exports")
    (tmp_path / "files").mkdir()
    dbmod.init()
    return tmp_path


def test_builtin_skills_are_valid():
    names = sorted(p.parent.name for p in config.BUILTIN_SKILLS.glob("*/SKILL.md"))
    assert names == ["deadlines-report", "document-report", "invoice-table"]
    for p in config.BUILTIN_SKILLS.glob("*/SKILL.md"):
        meta, body = skills.parse(p.read_text())
        assert meta["name"] == p.parent.name and body


@pytest.mark.parametrize("text,error", [
    ("no frontmatter", "frontmatter"),
    ("---\nname: Weekly Brief\ndescription: x\n---\nbody", "lowercase"),
    ("---\nname: weekly\n---\nbody", "description"),
])
def test_invalid_skills_are_rejected(text, error):
    with pytest.raises(ValueError, match=error):
        skills.parse(text)


def test_create_import_toggle_and_delete(env):
    skills.save("weekly-brief", "Summarize the week's documents.", "1. Call list_documents.")
    assert skills.get("weekly-brief").enabled and not skills.get("weekly-brief").builtin
    skills.set_enabled("weekly-brief", False)
    assert "weekly-brief" not in [s.name for s in skills.enabled()]

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("contracts/SKILL.md", SKILL.replace("weekly-brief", "contract-check"))
        z.writestr("contracts/reference/clauses.md", "Renewal clauses: ...")
        z.writestr("contracts/scripts/run.py", "print('never executed')")
        z.writestr("contracts/../../escape.md", "outside")
    skill = skills.import_file("contracts.zip", buf.getvalue())
    assert skill.references == ["reference/clauses.md"] and skill.ignored == ["scripts/run.py"]
    assert "Renewal" in skill.read_reference("reference/clauses.md")
    assert "No file" in skill.read_reference("../../escape.md") and not (env / "escape.md").exists()
    assert "No file" in skill.read_reference("scripts/run.py")

    assert skills.delete("contract-check") and skills.get("contract-check") is None
    assert not skills.delete("invoice-table")   # built-in skills can be disabled, not deleted


def _doc(name, ext, text, created="2026-10-20 10:00:00"):
    doc_id = ingest.add(name, ext, data=text.encode())
    with dbmod.db() as conn:
        conn.execute("UPDATE documents SET status='ready', pages_total=1, created_at=? WHERE id=?", (created, doc_id))
        conn.execute("INSERT INTO pages(doc_id, n, text, method) VALUES (?,1,?,'text')", (doc_id, text))
    return doc_id


def test_list_and_read_documents(env):
    from datetime import datetime, timedelta
    recent = (datetime.now() - timedelta(days=1)).isoformat()
    old = (datetime.now() - timedelta(days=40)).isoformat()
    _doc("Offerta.txt", ".txt", "Offerta 2027", created=recent.replace("T", " ")[:19])
    _doc("Vecchia.txt", ".txt", "Vecchia", created=old.replace("T", " ")[:19])
    inv = _doc("Invoice_118.pdf", ".pdf", "Invoice NW-118 total 4,000.00 EUR", created=recent.replace("T", " ")[:19])
    run = tools.Run()
    listing = tools.list_documents(run, days=7)
    assert "Offerta" in listing and "Invoice_118.pdf" in listing and "Vecchia" not in listing
    assert tools.list_documents(run, text="vecch").count("doc_id") == 1
    text = tools.read_document(run, inv)
    assert text.startswith("[1] Invoice_118.pdf, p. 1") and run.sources[0]["doc_id"] == inv
    tools.read_document(run, inv)
    assert len(run.sources) == 1                                     # the same page keeps its number
    assert "No document" in tools.read_document(run, "nope")


def test_create_files(env):
    run = tools.Run()
    tools.create_file(run, "Report: Q4/2026", "docx", content="# Report\n\nTotal **4,000 EUR**\n\n- one\n\n| A | B |\n|---|---|\n| 1 | 2 |")
    tools.create_file(run, "invoices", "xlsx", rows=[["Number", "Total"], ["NW-118", 4000]])
    assert [f["name"] for f in run.files] == ["Report- Q4-2026.docx", "invoices.xlsx"]
    doc = Docx(env / "exports" / run.id / "Report- Q4-2026.docx")
    assert doc.paragraphs[0].text == "Report" and doc.tables[0].cell(1, 1).text == "2"
    ws = load_workbook(env / "exports" / run.id / "invoices.xlsx").active
    assert ws["B2"].value == 4000
    assert "required" in tools.create_file(run, "x", "xlsx")


def test_agent_runs_tools_then_answers(env, monkeypatch):
    inv = _doc("Invoice_118.pdf", ".pdf", "Invoice NW-118 total 4,000.00 EUR, due 30/11/2026")
    replies = iter([
        {"content": "", "tool_calls": [{"id": "c1", "function": {"name": "read_document", "arguments": json.dumps({"doc_id": inv})}}]},
        {"content": "", "tool_calls": [{"id": "c2", "function": {"name": "create_file", "arguments": json.dumps(
            {"filename": "invoices", "format": "xlsx", "rows": [["Number", "Total"], ["NW-118", 4000]]})}}]},
        {"content": "Invoice NW-118: 4,000.00 EUR, due 30/11/2026 [1]."},
    ])
    seen = []

    async def fake_complete(role, messages, **params):
        seen.append(messages[-1])
        return next(replies)
    monkeypatch.setattr(llm, "complete", fake_complete)
    skill = skills.Skill("invoice-table", "Extract invoices", "Read invoices.", env, builtin=True)

    async def collect():
        return [e async for e in agent.run(skill, [{"role": "user", "content": "invoices to excel"}])]
    events = asyncio.run(collect())
    assert [e["type"] for e in events] == ["skill", "step", "step", "sources", "token", "file"]
    assert events[1] == {"type": "step", "tool": "read_document", "detail": inv}
    assert seen[1]["role"] == "tool" and "4,000.00 EUR" in seen[1]["content"]   # the tool result went back to the model
    assert events[3]["sources"][0]["name"] == "Invoice_118.pdf" and events[5]["name"] == "invoices.xlsx"


def test_router_picks_a_skill_or_none(env, monkeypatch):
    skills.save("weekly-brief", "Summarize the week's documents.", "...")

    async def fake_chat(role, messages, **params):
        assert role == "assistant" and "weekly-brief" in messages[0]["content"]
        assert params["response_format"]["json_schema"]["schema"]["properties"]["skill"]["enum"][-1] == "none"
        return json.dumps({"skill": "weekly-brief" if "week" in messages[1]["content"] else "none"})
    monkeypatch.setattr(llm, "chat", fake_chat)
    enabled = skills.enabled()
    assert asyncio.run(agent.route("brief me on this week", enabled)).name == "weekly-brief"
    assert asyncio.run(agent.route("who signed the contract?", enabled)) is None
