# /// script
# dependencies = ["pillow", "python-docx"]
# ///
"""English demo set for screenshots: interlinked office documents (scans + Word) in samples/demo/."""
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent.parent / "samples" / "demo"
FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"


def scan(name, lines, rotate=0.6):
    img = Image.new("RGB", (1240, 1754), "white")
    d = ImageDraw.Draw(img)
    y = 90
    for size, text in lines:
        if text == "---":
            d.line((80, y + 10, 1160, y + 10), fill="gray", width=2); y += 30; continue
        d.text((90, y), text, font=ImageFont.truetype(FONT, size), fill="black")
        y += int(size * 1.7)
    img.rotate(rotate, fillcolor="white").save(OUT / name)


def word(name, title, paragraphs):
    doc = Document()
    doc.add_heading(title, 1)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(OUT / name)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    scan("Invoice_NW-2026-118.png", [
        (44, "INVOICE NW-2026-118"), (28, "Date: 02/10/2026"),
        (28, "From: Northwind Components Ltd - VAT GB 284 7391 56"), (28, "221 Baker Road, Leeds LS1 4AP, United Kingdom"),
        (28, "To: Acme Studio S.r.l., Via Roma 12, Milano"), (20, " "), (28, "---"),
        (28, "Description                               Qty      Unit         Total"),
        (28, "---"), (28, "Sensor module SM-40                       40      85.00     3,400.00"),
        (28, "Calibration service                        1     600.00       600.00"), (28, "---"),
        (32, "Total due: 4,000.00 EUR"), (28, "Payment due by 30/11/2026 - IBAN GB33 BUKB 2020 1555 5555 55"),
    ])
    scan("Invoice_NW-2026-131.png", [
        (44, "INVOICE NW-2026-131"), (28, "Date: 21/10/2026"),
        (28, "From: Northwind Components Ltd - VAT GB 284 7391 56"),
        (28, "To: Acme Studio S.r.l., Via Roma 12, Milano"), (28, "---"),
        (28, "Harbor website - hardware kiosk units      2   1,250.00     2,500.00"), (28, "---"),
        (32, "Total due: 2,500.00 EUR"), (28, "Payment due by 20/12/2026"),
    ], rotate=-0.8)
    word("Service_Agreement_Northwind.docx", "Service Agreement", [
        "This Service Agreement is made on 15 January 2026 between Acme Studio S.r.l. (VAT IT 09876543210), "
        "Via Roma 12, Milano, and Northwind Components Ltd (VAT GB 284 7391 56), 221 Baker Road, Leeds.",
        "Northwind Components Ltd will supply sensor modules and calibration services to Acme Studio S.r.l. "
        "for a period of 12 months.",
        "The agreement expires on 31/12/2026 and renews automatically unless cancelled in writing by 30/11/2026.",
        "Payment terms: invoices are due within 60 days.",
        "Signed for Acme Studio S.r.l.: Luca Moretti, Operations Manager (luca.moretti@acmestudio.it).",
        "Signed for Northwind Components Ltd: Sarah Collins, Chief Financial Officer (s.collins@northwind.co.uk).",
    ])
    word("Minutes_Harbor_Kickoff.docx", "Meeting minutes - Harbor Website Redesign kickoff", [
        "Date: 08/10/2026. Attendees: Luca Moretti (Acme Studio S.r.l.), Giulia Ferri, Project Lead (Acme Studio S.r.l.), "
        "and Tom Becker from Harbor Logistics GmbH.",
        "Harbor Logistics GmbH commissioned the Harbor Website Redesign project from Acme Studio S.r.l.; "
        "Giulia Ferri leads the project.",
        "Decisions: the new website goes live on 15/01/2027. The design review with Harbor Logistics GmbH is due by 20/11/2026.",
        "Two kiosk units for the reception will be purchased from Northwind Components Ltd.",
        "Action: Giulia Ferri sends the Q3 progress report by 15/10/2026.",
    ])
    word("Purchase_Order_PO-7781.docx", "Purchase order PO-7781", [
        "Date: 25/10/2026. Buyer: Acme Studio S.r.l., Via Roma 12, Milano.",
        "Supplier: Northwind Ltd, VAT GB 284 7391 65.",
        "Items: 20 x sensor module SM-40 at 85.00 EUR. Requested delivery by 10/11/2026.",
        "Approved by Luca Moretti.",
    ])
    word("Remote_Work_Policy.docx", "Remote work policy 2027", [
        "Applies to all employees of Acme Studio S.r.l.",
        "Employees may work remotely up to 3 days per week. Requests for the 2027 schedule must be submitted "
        "to HR by 15/12/2026.",
        "Contact: Giulia Ferri for the design team, Luca Moretti for operations.",
    ])
    print("demo files:", sorted(p.name for p in OUT.iterdir()))


if __name__ == "__main__":
    main()
